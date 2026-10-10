"""Load the LBNL interconnection queue workbook into queue_projects.

The workbook's exact column names can change between editions, so columns are found by
alias and the header row is detected instead of assumed.
"""
import os
import re

import openpyxl

from ..config import Settings
from .common import download_first_available, normalize_county

COLUMN_ALIASES = {
    "queue_id": ("q_id", "queue_id", "queue id", "project_id"),
    "status": ("q_status", "status", "queue_status", "queue status"),
    "state": ("state", "state_abbr", "state_code"),
    "county": ("county", "counties"),
    "tech": ("type_clean", "type", "fuel", "resource_type", "technology", "type1"),
    "mw": ("mw1", "mw", "capacity_mw", "mw_1", "capacity (mw)"),
}
REQUIRED_COLUMNS = ("status", "state", "county", "tech", "mw")


def find_columns(header_cells: list) -> dict:
    """Map our field names to column positions using the aliases. Missing fields are omitted."""
    lowered = [str(cell).strip().lower() if cell is not None else "" for cell in header_cells]
    found = {}
    for field, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            if alias in lowered:
                found[field] = lowered.index(alias)
                break
    return found


def locate_table(workbook) -> tuple:
    """Return (sheet, header_row_index, columns) for the sheet that best matches the aliases."""
    best = None
    for sheet in workbook.worksheets:
        for row_index, row in enumerate(sheet.iter_rows(min_row=1, max_row=25, values_only=True)):
            columns = find_columns(list(row))
            matched = sum(1 for field in REQUIRED_COLUMNS if field in columns)
            if best is None or matched > best[0]:
                best = (matched, sheet, row_index, columns)
    if best is None or best[0] < len(REQUIRED_COLUMNS):
        seen = {sheet.title: [c for c in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), [])]
                for sheet in workbook.worksheets}
        raise RuntimeError(
            "Could not find the queue table. Needed columns: "
            f"{REQUIRED_COLUMNS}. Sheets and first rows seen: {seen}. "
            "Run 'python -m ranker.ingest inspect-queue' and adjust COLUMN_ALIASES."
        )
    _, sheet, row_index, columns = best
    return sheet, row_index, columns


def normalize_status(value) -> str:
    text = str(value or "").lower()
    if "withdraw" in text:
        return "withdrawn"
    if "operation" in text:
        return "operational"
    if "active" in text or "suspend" in text:
        return "active"
    return "other"


def normalize_tech(value) -> str:
    text = str(value or "").lower()
    if "storage" in text or "battery" in text:
        return "storage"
    if "solar" in text or "photovoltaic" in text:
        return "solar"
    if "wind" in text:
        return "wind"
    if "gas" in text:
        return "gas"
    return "other"


def split_counties(value) -> list:
    """Projects can list several counties ("Harris, Fort Bend"). Returns normalised names."""
    parts = re.split(r"[;,/&]| and ", str(value or ""))
    names = [normalize_county(part) for part in parts]
    return [name for name in names if name]


def parse_rows(workbook, state: str) -> list:
    """Return (queue_id, county_norm, status, tech, mw) rows for one state.

    A project in several counties has its MW split evenly across them.
    """
    sheet, header_index, columns = locate_table(workbook)
    rows = []
    for row in sheet.iter_rows(min_row=header_index + 2, values_only=True):
        if row[columns["state"]] is None or str(row[columns["state"]]).strip().upper() != state:
            continue
        try:
            megawatts = float(row[columns["mw"]])
        except (TypeError, ValueError):
            continue
        if megawatts <= 0:
            continue
        counties = split_counties(row[columns["county"]])
        if not counties:
            continue
        queue_id = str(row[columns["queue_id"]]) if "queue_id" in columns and row[columns["queue_id"]] else None
        share = megawatts / len(counties)
        for county in counties:
            rows.append((queue_id, county, normalize_status(row[columns["status"]]),
                         normalize_tech(row[columns["tech"]]), share))
    return rows


def ingest_queue(connection, settings: Settings) -> int:
    path = download_first_available([settings.queue_url], os.path.join(settings.data_dir, "queue.xlsx"))
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = parse_rows(workbook, settings.state)
    with connection.cursor() as cursor:
        cursor.execute("TRUNCATE queue_projects")
        cursor.executemany(
            "INSERT INTO queue_projects (queue_id, county_norm, status, tech, mw) VALUES (%s, %s, %s, %s, %s)",
            rows,
        )
    connection.commit()
    print(f"Loaded {len(rows)} queue rows for {settings.state}")
    return len(rows)


def inspect_queue(settings: Settings) -> None:
    """Print each sheet's name and first rows so column names can be checked."""
    path = download_first_available([settings.queue_url], os.path.join(settings.data_dir, "queue.xlsx"))
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        print(f"\n== {sheet.title}")
        for row in sheet.iter_rows(min_row=1, max_row=6, values_only=True):
            print([cell for cell in row][:20])
